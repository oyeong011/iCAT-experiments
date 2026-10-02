#!/usr/bin/env python3
"""장시간 실행에 따른 누적 WAF (FIO-Fast, 단일 워크로드): iCAT-v4 10시간 vs 고정 CAT (기본 37 / 견고 50)
-> figs/longrun_cumulative.(png|svg|pdf), figs/data/longrun_cumulative.xlsx (데이터 + 엑셀 차트, 직접 편집용)
원천: result/mix-20260911/<run>/control-series.txt (30초 누적 카운터). v4 = t4-onlinev4-rep1-x60 한 실행의 1·3·10시간 시점 값.
고정값: CAT-50 = t4-fixed50-rep1-x18(3시간), CAT-37 = t4-fixed37-rep1-x18(3시간, 없으면 10분 실행 t4-fixed37-rep1). CAT-47 = 시트에만(그림에는 안 그림)."""
import re
from pathlib import Path
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.styles import Font
font_manager.fontManager.addfont('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
M = Path('/home/oy/iCAT/result/mix-20260911')
INK, INK2, V4 = '#1d1d1f', '#6e6e73', '#2a78d6'
plt.rcParams.update({'font.family': 'Noto Sans CJK JP', 'font.size': 10.5, 'axes.facecolor': 'white', 'figure.facecolor': 'white',
                     'axes.spines.top': False, 'axes.spines.right': False, 'axes.spines.left': False, 'axes.edgecolor': '#d2d2d7',
                     'axes.grid': True, 'axes.grid.axis': 'y', 'grid.color': '#ededf0', 'grid.linewidth': 0.8, 'axes.axisbelow': True,
                     'xtick.color': INK2, 'ytick.color': INK2, 'ytick.left': False, 'axes.labelcolor': INK2, 'svg.fonttype': 'path'})
def cum(d):
    s = [tuple(map(int, m.groups())) for l in open(M / d / 'control-series.txt') if (m := re.match(r'(\d+) .*active=1 .*host_pages=(\d+) gc_pages=(\d+)', l))]
    return [((t - s[0][0]) / 3600, 1 + g / h) for t, h, g in s[1:] if h]
c37 = 't4-fixed37-rep1-x18' if (M / 't4-fixed37-rep1-x18' / 'summary.txt').exists() else 't4-fixed37-rep1'
RUNS = [('iCAT-v4 (제안)', 't4-onlinev4-rep1-x60', V4, 2.4, True), ('CAT-50 (견고 설정)', 't4-fixed50-rep1-x18', '#8e8e93', 1.8, True),
        ('CAT-37 (기본값)', c37, '#e3a33b', 1.8, True), ('CAT-47 (최적)', 't4-fixed47-rep1-x18', '#1d1d1f', 1.0, False)]
D = {lab: cum(d) for lab, d, *_ in RUNS}
at = lambda c, h: min(c, key=lambda p: abs(p[0] - h))[1] if c and c[-1][0] >= h - 0.02 else None
v = D['iCAT-v4 (제안)']
# ---- PNG
fig, ax = plt.subplots(figsize=(11, 5.6))
for lab, d, col, lw, draw in RUNS:
    if not draw: continue
    c = D[lab]
    if c[-1][0] < 0.5:   # only a 10-min run so far: show its value as a dashed reference, clearly labelled (3 h run queued)
        ax.axhline(c[-1][1], color=col, lw=1.6, ls=(0, (5, 3)), label=f'{lab}  {c[-1][1]:.3f} (10분 실행값, 3시간 측정 예정)'); continue
    ax.plot([h for h, _ in c], [w for _, w in c], color=col, lw=lw, label=lab + ('' if c[-1][0] > 9 else f'  {c[-1][1]:.3f} (3시간 측정)'), zorder=3 if lab.startswith('iCAT') else 2)
for h in (1, 3, 10):
    w = at(v, h); ax.scatter([h], [w], s=40, color=V4, zorder=4, edgecolor='white', linewidth=1.2)
    ax.annotate(f'{h}시간: {w:.3f}', (h, w), (18 if h == 1 else 0, 16), textcoords='offset points', ha='left' if h == 1 else ('right' if h == 10 else 'center'), fontsize=10, color=V4, weight='bold')
ax.set_xlim(0, 10.3); ax.set_xticks(range(0, 11)); ax.set_ylim(1.75, 2.05)
ax.set_xlabel('실행 시간 (시간)'); ax.set_ylabel('누적 WAF (낮을수록 좋음)')
ax.legend(loc='upper right', fontsize=10, frameon=False)
fig.suptitle('FIO-Fast 장시간 실행에서의 누적 WAF', x=0.06, ha='left', y=0.99, fontsize=13.5, weight='bold', color=INK)
fig.text(0.06, 0.925, 'iCAT-v4는 10시간 실행 한 번의 누적값(30초 기록). 고정 CAT는 별도 실행이며, 고정값은 파라미터가 바뀌지 않아 약 20분 이후 거의 일정하다.', fontsize=9.5, color=INK2)
fig.tight_layout(rect=(0, 0, 1, 0.91))
for e in ('png', 'svg', 'pdf'): fig.savefig(f'/home/oy/iCAT/figs/longrun_cumulative.{e}', dpi=220, bbox_inches='tight')
# ---- xlsx
wb = Workbook(); ws = wb.active; ws.title = '그림6 데이터'
col = 1; refs = {}
for lab, d, c, lw, draw in RUNS:
    ws.cell(row=1, column=col, value=f'{lab} 시간(h)').font = Font(bold=True); ws.cell(row=1, column=col + 1, value=f'{lab} 누적 WAF').font = Font(bold=True)
    for i, (h, w) in enumerate(D[lab], 2): ws.cell(row=i, column=col, value=round(h, 4)); ws.cell(row=i, column=col + 1, value=round(w, 4))
    refs[lab] = (col, len(D[lab]), c, lw, draw); col += 2
ws.cell(row=1, column=col + 1, value='iCAT-v4 시점별 누적 WAF').font = Font(bold=True)
for i, h in enumerate((1, 3, 10), 2): ws.cell(row=i, column=col + 1, value=f'{h}시간'); ws.cell(row=i, column=col + 2, value=round(at(v, h), 3))
ws.cell(row=6, column=col + 1, value='고정 CAT 최종 누적 WAF').font = Font(bold=True)
for i, (lab, d, *_) in enumerate(RUNS[1:], 7): ws.cell(row=i, column=col + 1, value=f'{lab} ({D[lab][-1][0]:.1f}시간)'); ws.cell(row=i, column=col + 2, value=round(D[lab][-1][1], 3))
ws.cell(row=11, column=col + 1, value='※ CAT-47은 데이터만 있고 차트에는 그리지 않음. 넣으려면 차트 → 데이터 선택 → 추가.')
ch = ScatterChart(); ch.title = 'FIO-Fast 장시간 실행에서의 누적 WAF (낮을수록 좋음)'; ch.style = 13; ch.scatterStyle = 'lineMarker'; ch.display_blanks = 'span'
ch.x_axis.title = '실행 시간 (시간)'; ch.y_axis.title = '누적 WAF'; ch.height, ch.width = 12, 26; ch.legend.position = 'b'
ch.x_axis.delete = False; ch.y_axis.delete = False; ch.x_axis.scaling.min = 0; ch.x_axis.scaling.max = 10.5; ch.x_axis.majorUnit = 1
ch.y_axis.scaling.min = 1.7; ch.y_axis.scaling.max = 2.1; ch.y_axis.number_format = '0.00'
for lab, (c0, n, c, lw, draw) in refs.items():
    if not draw: continue
    s = Series(Reference(ws, min_col=c0 + 1, min_row=1, max_row=n + 1), Reference(ws, min_col=c0, min_row=2, max_row=n + 1), title_from_data=True)
    s.smooth = False; s.marker.symbol = 'none'; s.graphicalProperties.line.solidFill = c[1:].upper(); s.graphicalProperties.line.width = int(lw * 12700)
    ch.series.append(s)
ws.add_chart(ch, f'{ws.cell(row=1, column=col + 5).column_letter}2')
wb.save('/home/oy/iCAT/figs/data/longrun_cumulative.xlsx')
print('CAT-37 from', c37, '| v4 1/3/10h', [round(at(v, h), 3) for h in (1, 3, 10)], '| fixed end', {k: round(D[k][-1][1], 3) for k in D if not k.startswith('iCAT')})
