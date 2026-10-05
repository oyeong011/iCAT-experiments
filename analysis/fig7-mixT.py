#!/usr/bin/env python3
"""그림 7 (불규칙 순서 혼합 10시간, mixT). 9구간: YCSB-A, OLTP, Varmail, YCSB-A, YCSB-B, FIO-Fast, Varmail, YCSB-A, OLTP.
(가) 10분 구간 WAF: iCAT-v4 vs 고정 CAT (있는 것만), (나) v4 학습 상태, (다) v4 활성 후보 수 (파티션 2), 재설정 표시.
-> figs/fig7_mixT.(png|svg|pdf), figs/data/fig7_mixT.xlsx (데이터 + 엑셀 차트)
원천: 각 실행의 control-series.txt, phase-X-start.time; v4 학습 기록은 kernel 로그(analysis/v4-log.py)."""
import re, datetime as dt, importlib.util
from pathlib import Path
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import Patch
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.styles import Font
font_manager.fontManager.addfont('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
spec = importlib.util.spec_from_file_location('v4log', '/home/oy/iCAT/analysis/v4-log.py'); v4log = importlib.util.module_from_spec(spec); spec.loader.exec_module(v4log)
R = Path('/home/oy/iCAT/result'); M = R / 'mix-20260911'
def pick(p, other):   # this PC first, else the second PC's copy
    d = M / f'mixT-{p}-rep1'
    return d if (d / 'summary.txt').exists() else R / other / f'mixT-{p}-rep1'
POL = [('iCAT-v4 (제안)', M / 'mixT-onlinev4-rep1', '#2a78d6', 2.4), ('CAT-50 (견고 설정)', pick('fixed50', 'mixT-fixed50-tenhour-20261004'), '#8e8e93', 1.6),
       ('CAT-37 (기본값)', pick('fixed37', 'mixT-fixed37-tenhour-20261004'), '#e3a33b', 1.6)]
PART, STEP = '2', 10 / 60
NAMES = ['YCSB-A', 'OLTP', 'Varmail', 'YCSB-A', 'YCSB-B', 'FIO-Fast', 'Varmail', 'YCSB-A', 'OLTP']
BAND = {'YCSB-A': '#f6f1e7', 'YCSB-B': '#f3ece0', 'OLTP': '#eaf3ee', 'Varmail': '#eef2fb', 'FIO-Fast': '#f7ecec'}
INK, INK2, V4, PUR = '#1d1d1f', '#6e6e73', '#2a78d6', '#7d3cff'
STATE = [('전수 순회', '#d9d9de'), ('탐색', '#8e8e93'), ('정착', V4)]
plt.rcParams.update({'font.family': 'Noto Sans CJK JP', 'font.size': 10.5, 'axes.facecolor': 'white', 'figure.facecolor': 'white',
                     'axes.spines.top': False, 'axes.spines.right': False, 'axes.spines.left': False, 'axes.edgecolor': '#d2d2d7',
                     'axes.grid': True, 'axes.grid.axis': 'y', 'grid.color': '#ededf0', 'grid.linewidth': 0.8, 'axes.axisbelow': True,
                     'xtick.color': INK2, 'ytick.color': INK2, 'ytick.left': False, 'axes.labelcolor': INK2, 'svg.fonttype': 'path'})

def run(d):   # 10-min bin WAF on the run's own clock, phase starts (h), totals
    raw = [tuple(map(int, m.groups())) for l in open(d / 'control-series.txt') if (m := re.match(r'(\d+) .*active=1 .*host_pages=(\d+) gc_pages=(\d+)', l))]
    t0 = raw[0][0]; s = [((t - t0) / 3600, h, g) for t, h, g in raw]
    out, k = [], 0.0
    while k + STEP <= s[-1][0] + 1e-9:
        a = min(s, key=lambda p: abs(p[0] - k)); b = min(s, key=lambda p: abs(p[0] - k - STEP))
        if b[1] > a[1]: out.append((round(k + STEP / 2, 4), 1 + (b[2] - a[2]) / (b[1] - a[1])))
        k += STEP
    starts = [(dt.datetime.fromisoformat((d / f'phase-{c}-start.time').read_text().strip()).timestamp() - t0) / 3600 for c in 'ABCDEFGHI']
    sm = (d / 'summary.txt').read_text()
    tot = float(re.search(r'^total .*WAF=([\d.]+)', sm, re.M)[1])
    ph = [float(re.search(rf'^phase{c}\(\S+\) .*WAF=([\d.]+)', sm, re.M)[1]) for c in 'ABCDEFGHI']
    return out, starts, tot, ph, s[-1][0]
pol = [(lab, d, c, lw) for lab, d, c, lw in POL if (d / 'summary.txt').exists()]
D = {lab: run(d) for lab, d, *_ in pol}
_, starts, _, _, X1 = D['iCAT-v4 (제안)']
bounds = starts + [X1]

# v4 learner trace, partition PART
L = v4log.lines(str(M / 'mixT-onlinev4-rep1'))
ts = lambda l: float(re.match(r'\[\s*(\d+\.\d+)\]', l)[1])
t0 = next(ts(l) for l in L if 'Measurement: command=start' in l)
tr = [((ts(l) - t0) / 3600, int(m[1]), int(m[2]), int(m[3])) for l in L if (m := re.search(rf'part={PART} epoch=(\d+) window=\d+ .*?settled=(\d) active=(\d+)', l))]
resets = [(ts(l) - t0) / 3600 for l in L if 'WATGC_V2 reset' in l and f'part={PART} ' in l]
stt, swept = [], False
for i, (h, ep, st_, act) in enumerate(tr):
    if i and ep != tr[i - 1][1]: swept = False
    elif i and act < tr[i - 1][3]: swept = True
    stt.append(2 if st_ else (1 if swept else 0))
def bin_last(vals):   # per 10-min bin: last value (candidates) / longest-held state
    out, k = [], 0.0
    while k + STEP <= X1 + 1e-9:
        idx = [i for i, r in enumerate(tr) if k <= r[0] < k + STEP]
        if idx: out.append((round(k + STEP / 2, 4), vals(idx)))
        k += STEP
    return out
cand = bin_last(lambda idx: tr[idx[-1]][3])
state = bin_last(lambda idx: max(set(stt[i] for i in idx), key=[stt[i] for i in idx].count))

# ---- PNG
fig, ax = plt.subplots(3, 1, figsize=(14, 8.2), sharex=True, gridspec_kw={'height_ratios': [1.35, 0.16, 0.9], 'hspace': 0.3})
for i, n in enumerate(NAMES):
    for a_ in (ax[0], ax[2]): a_.axvspan(bounds[i], bounds[i + 1], color=BAND[n], lw=0, zorder=0)
    ax[0].text((bounds[i] + bounds[i + 1]) / 2, 1.01, f'{i + 1}. {n}', transform=ax[0].get_xaxis_transform(), ha='center', va='bottom', fontsize=9, color=INK, weight='bold')
    for a_ in ax: a_.axvline(bounds[i], color='#c7c7cc', lw=0.8)
for lab, d, c, lw in pol:
    w, _, tot, ph, _ = D[lab]
    ax[0].plot([h for h, _ in w], [v for _, v in w], color=c, lw=lw, marker='o', ms=3, label=f'{lab}   10시간 전체 WAF {tot:.3f}', zorder=3 if 'v4' in lab else 2)
ax[0].set_ylabel('구간 WAF (10분 단위)\n(낮을수록 좋음)'); ax[0].legend(loc='upper left', fontsize=9.5, frameon=True, facecolor='white', edgecolor='#e0dfdb', framealpha=0.95)
ax[0].set_title('(가) 10분 구간 WAF', loc='left', fontsize=11, color=INK, weight='bold', pad=22)
for h, s_ in state: ax[1].axvspan(h - STEP / 2, h + STEP / 2, color=STATE[s_][1], lw=0)
ax[1].set_yticks([]); ax[1].grid(False); ax[1].set_title(f'(나) v4 학습 상태 (파티션 {PART}, 10분마다 가장 오래 머문 상태)', loc='left', fontsize=10.5, color=INK2)
ax[1].legend(handles=[Patch(color=c, label=n) for n, c in STATE], loc='lower right', bbox_to_anchor=(1.0, 1.0), ncol=3, fontsize=9, frameon=False, handlelength=1.0)
ax[2].plot([h for h, _ in cand], [v for _, v in cand], color=V4, lw=2.2, drawstyle='steps-mid')
chg = [(h, v) for i, (h, v) in enumerate(cand) if i == 0 or v != cand[i - 1][1]]
ax[2].scatter([h for h, _ in chg], [v for _, v in chg], s=24, color=V4, zorder=4)
for i, (h, v) in enumerate(chg):
    if i == 0 or abs(v - chg[i - 1][1]) >= 4: ax[2].annotate(str(v), (h, v), (2, 7), textcoords='offset points', fontsize=9.5, color=V4, weight='bold')
for r in resets:
    for a_ in ax: a_.axvline(r, color=PUR, lw=1.6)
    ax[2].annotate('변화 감지\n→ 재설정', (r, 46), (r + 0.08, 46), fontsize=9, color=PUR, weight='bold', va='top')
# explain the two odd stretches (checked in the kernel log): FIO phase windows all discarded; YCSB-A collapse
f0, f1 = bounds[5], bounds[6]
ax[1].text((f0 + f1) / 2, 0.5, '판단 구간 없음', transform=ax[1].get_xaxis_transform(), ha='center', va='center', fontsize=9, color=INK2)
ax[2].text((f0 + f1) / 2, 30, 'GC가 너무 많아\n판단 구간이 모두\n폐기됨(327개)\n→ 학습 없음', ha='center', va='center', fontsize=9, color=INK2)
lo_h = min((h for h, v in cand if bounds[3] <= h <= bounds[4]), key=lambda h: dict(cand)[h])
ax[2].annotate('4번 YCSB-A: GC가 거의 없어 쓰던 조합의 WAF≈1.00\n→ 다른 후보가 모두 "15% 이상 나쁨"으로 제거됨', (lo_h, 1), (bounds[2] + 0.1, 24), fontsize=8.5, color=INK2, va='center',
               bbox=dict(boxstyle='round,pad=0.3', fc='white', ec='#d2d2d7', lw=0.8), arrowprops=dict(arrowstyle='->', color=INK2, lw=0.9))
ax[2].set_ylim(0, 52); ax[2].set_yticks([0, 15, 30, 45]); ax[2].set_ylabel('남은 후보 수'); ax[2].set_xlabel('측정 시작 후 시간 (시간)')
ax[2].set_title(f'(다) v4 활성 후보 수 (파티션 {PART}, 바뀌는 지점에 점)', loc='left', fontsize=10.5, color=INK2)
ax[2].set_xlim(0, X1); ax[2].set_xticks(range(0, int(X1) + 1))
fig.suptitle('그림 7. 불규칙 순서 응용 혼합 10시간 실행에서 iCAT-v4의 WAF, 학습 상태, 활성 후보 수', x=0.06, ha='left', y=0.995, fontsize=13.5, weight='bold', color=INK)
fig.text(0.06, 0.957, '9구간 각 4000초. 배경색 = 워크로드. 보라 선 = v4가 워크로드 변화를 감지해 재설정한 시점(전환 8번 중 2번). 고정 CAT은 같은 조건의 별도 실행(두 번째 PC), 각 1회.', fontsize=9.5, color=INK2)
for e in ('png', 'svg', 'pdf'): fig.savefig(f'/home/oy/iCAT/figs/fig7_mixT.{e}', dpi=200, bbox_inches='tight')

# ---- xlsx
wb = Workbook(); ws = wb.active; ws.title = '그림7 불규칙 혼합 10시간'
cols = []
for lab, *_ in pol: w = D[lab][0]; cols += [(f'{lab} 시간(h)', [h for h, _ in w]), (f'{lab} 10분 구간 WAF', [round(v, 4) for _, v in w])]
cols += [('시간(h)', [h for h, _ in state]), (f'v4 학습 상태 p{PART} (0=순회,1=탐색,2=정착)', [v for _, v in state])]
cols += [('시간(h) ', [h for h, _ in cand]), (f'v4 남은 후보 수 p{PART}', [v for _, v in cand])]
for j, (h, v) in enumerate(cols, 1):
    ws.cell(row=1, column=j, value=h).font = Font(bold=True)
    for i, val in enumerate(v, 2): ws.cell(row=i, column=j, value=val)
c0 = len(cols) + 2
for j, h in enumerate(['구간', '워크로드', '시작(h, v4)'] + [lab for lab, *_ in pol], c0): ws.cell(row=1, column=j, value=h).font = Font(bold=True)
for i, n in enumerate(NAMES):
    ws.cell(row=2 + i, column=c0, value=i + 1); ws.cell(row=2 + i, column=c0 + 1, value=n); ws.cell(row=2 + i, column=c0 + 2, value=round(starts[i], 3))
    for k, (lab, *_) in enumerate(pol): ws.cell(row=2 + i, column=c0 + 3 + k, value=round(D[lab][3][i], 3))
ws.cell(row=11, column=c0 + 1, value='10시간 전체').font = Font(bold=True)
for k, (lab, *_) in enumerate(pol): ws.cell(row=11, column=c0 + 3 + k, value=round(D[lab][2], 3))
ws.cell(row=13, column=c0, value='재설정(h, p' + PART + ')').font = Font(bold=True); ws.cell(row=13, column=c0 + 1, value=', '.join(f'{r:.3f}' for r in resets))
L_ = ws.cell(row=1, column=c0 + 8).column_letter
def chart(title, yt, ser, ymin, ymax, anchor):
    ch = ScatterChart(); ch.title = title; ch.style = 13; ch.scatterStyle = 'lineMarker'; ch.display_blanks = 'span'
    ch.x_axis.title = '측정 시작 후 시간 (시간)'; ch.y_axis.title = yt; ch.height, ch.width = 10, 30; ch.legend.position = 'b'
    ch.x_axis.delete = False; ch.y_axis.delete = False; ch.x_axis.scaling.min = 0; ch.x_axis.scaling.max = round(X1 + 0.2); ch.x_axis.majorUnit = 1
    ch.y_axis.scaling.min = ymin; ch.y_axis.scaling.max = ymax
    for xc, color, width in ser:
        n = len(cols[xc][1]); s = Series(Reference(ws, min_col=xc + 2, min_row=1, max_row=n + 1), Reference(ws, min_col=xc + 1, min_row=2, max_row=n + 1), title_from_data=True)
        s.smooth = False; s.graphicalProperties.line.solidFill = color; s.graphicalProperties.line.width = width
        s.marker.symbol = 'circle'; s.marker.size = 3; s.marker.graphicalProperties.solidFill = color; s.marker.graphicalProperties.line.solidFill = color
        ch.series.append(s)
    ws.add_chart(ch, anchor)
wafs = [v for lab in D for _, v in D[lab][0]]; k = 2 * len(pol)
chart('(가) 불규칙 순서 응용 혼합 10시간: 10분 구간 WAF (낮을수록 좋음)', 'WAF', [(2 * i, c[1:].upper(), 22225) for i, (lab, d, c, lw) in enumerate(pol)], round(min(wafs) - 0.1, 1), round(max(wafs) + 0.1, 1), f'{L_}2')
chart('(나) v4 학습 상태 (0 = 전수 순회, 1 = 탐색, 2 = 정착)', '상태', [(k, '2A78D6', 22225)], -0.2, 2.2, f'{L_}23')
chart(f'(다) v4 활성 후보 수 (파티션 {PART})', '남은 후보 수', [(k + 2, '2A78D6', 28575)], 0, 48, f'{L_}44')
wb.save('/home/oy/iCAT/figs/data/fig7_mixT.xlsx')
print('drawn:', {l: round(D[l][2], 3) for l in D}, '| resets h', [round(r, 2) for r in resets])
