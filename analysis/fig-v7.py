#!/usr/bin/env python3
"""iCAT-v7 그림 2개 + 엑셀 (논문용 진한 색, 선 2pt).
그림 6: WAF 추이 - (가) 90분 혼합 FIO→Varmail→OLTP: CAT-37(기본)/CAT-50(견고)/CAT-47(최적)/v7, 5분 단위
                    (나) mixY 10시간 F/S/V 20구간: CAT-47/v6/v7, 15분 단위
그림 7: v7 학습 - 파티션 2의 남은 후보 수(45개에서 시작, 버리면 줄고, 변화 감지 시 다시 45개), 30초 스냅샷마다
-> figs/fig6_v7.(png|pdf), figs/fig7_v7.(png|pdf), figs/data/fig6_v7.xlsx, figs/data/fig7_v7.xlsx
원천: control-series.txt, phase-SNN-start.time, summary.txt, kernel-snapshots/NNNNNN.log (스냅샷 N = control-series N번째 줄). 모두 이 PC, 각 1회."""
import re, glob, datetime as dt
from pathlib import Path
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Font, PatternFill, Alignment
font_manager.fontManager.addfont('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
R = Path('/home/oy/iCAT/result'); FIG = Path('/home/oy/iCAT/figs'); C60 = R / 'cat60-90min-thispc-20261011'
V7A, V7B = R / 'v7-20261011/fvo90/mixY-onlinev7-rep1', R / 'v7-20261011/mixY/mixY-onlinev7-rep1'
# (label, run dir, color, line width) - v7 last so it sits on top
A = [('CAT-37 (기본)', C60 / 'mixY-arm37-rep1', '#8B2E16', 2), ('CAT-50 (견고)', C60 / 'mixY-arm50-rep1', '#5A5A5A', 2),
     ('CAT-47 (최적)', C60 / 'mixY-arm47-rep1', '#1A1A1A', 2), ('iCAT-v7 (제안)', V7A, '#0B3C8C', 2.8)]
B = [('CAT-47 (최적)', R / 'gc-20261010/mixY-fixed47-rep1', '#1A1A1A', 2), ('iCAT-v6', R / 'gc-20261010/mixY-onlinev6-rep1', '#9A6A00', 2),
     ('iCAT-v7 (제안)', V7B, '#0B3C8C', 2.8)]
WL = {'F': 'FIO-Fast', 'S': 'FIO-Slow', 'V': 'Varmail', 'O': 'OLTP'}
SEQA, SEQB = 'F V O'.split(), open(R / 'gc-20261010/mixy-seq.txt').read().split()
BAND = {'FIO-Fast': '#f7ecec', 'FIO-Slow': '#f6f1e7', 'Varmail': '#eef2fb', 'OLTP': '#eaf3ee'}
INK, INK2, NAVY, DPUR = '#1d1d1f', '#6e6e73', '#0B3C8C', '#4B1C7A'
NOTE = '정책마다 1회 실행, 모두 이 PC. 1% 안팎의 차이는 실행 간 오차 범위.'
plt.rcParams.update({'font.family': 'Noto Sans CJK JP', 'font.size': 10.5, 'axes.spines.top': False, 'axes.spines.right': False,
                     'axes.grid': True, 'axes.grid.axis': 'y', 'grid.color': '#ededf0', 'axes.axisbelow': True,
                     'xtick.color': INK2, 'ytick.color': INK2, 'axes.labelcolor': INK2})

def iso(p): return dt.datetime.fromisoformat(open(p).read().strip()).timestamp()
def rows(d):   # (unix time, host_pages, gc_pages) for every control-series line
    out = []
    for l in open(d / 'control-series.txt'):
        f = dict(kv.split('=') for kv in l.split()[1:] if '=' in kv)
        out.append((int(l.split()[0]), int(f['host_pages']), int(f['gc_pages'])))
    return out
def starts(d): return [iso(p) for p in sorted(glob.glob(f'{d}/phase-S*-start.time'))]
def binned(d, minutes, unit):   # WAF of each `minutes` bin on the run's own clock, x = bin middle in `unit` minutes
    rs, t0 = rows(d), starts(d)[0]; end = iso(sorted(glob.glob(f'{d}/phase-S*-end.time'))[-1]); out = []
    at = lambda t: min(rs, key=lambda r: abs(r[0] - t))
    b = t0
    while b + minutes * 60 <= end + 30:
        a, z = at(b), at(b + minutes * 60)
        if z[1] > a[1]: out.append(((b - t0) / 60 / unit + minutes / 2 / unit, 1 + (z[2] - a[2]) / (z[1] - a[1])))
        b += minutes * 60
    return out
def total(d): return float(re.search(r'WAF=([\d.]+)', next(l for l in open(d / 'summary.txt') if l.startswith('total')))[1])
def phase_waf(d, n): return [float(re.search(r'WAF=([\d.]+)', l)[1]) for l in open(d / 'summary.txt') if l.startswith('phaseS')][:n]
SAMPLE = re.compile(r'WATGC_V2 sample ns=\d+ part=(\d+) .*?best=(\d+) next=\d+ settled=\d+ active=(\d+)')
def candidates(d, unit, part='2'):   # last sample of `part` in each 30 s snapshot -> (time, remaining candidates, best arm)
    rs, t0, out = rows(d), starts(d)[0], []
    for f in sorted(glob.glob(f'{d}/kernel-snapshots/*.log')):
        i = int(Path(f).stem)
        last = [m for m in SAMPLE.finditer(open(f, errors='replace').read()) if m[1] == part]
        if i < len(rs) and last and rs[i][0] >= t0: out.append(((rs[i][0] - t0) / 60 / unit, int(last[-1][3]), int(last[-1][2])))
    return out

PANELS = [('(가) 90분 혼합: FIO-Fast → Varmail → OLTP (각 30분), 5분 단위', A, SEQA, 5, 1, '측정 시작 후 시간 (분)'),
          ('(나) mixY 10시간: FIO-Fast·FIO-Slow·Varmail 30분 × 20구간, 15분 단위', B, SEQB, 15, 60, '측정 시작 후 시간 (시간)')]
def bands(a, d, seq, unit, labels=True):
    st = [(s - starts(d)[0]) / 60 / unit for s in starts(d)]; ed = st[1:] + [st[-1] + 30 / unit]
    for i, (x0, x1) in enumerate(zip(st, ed)):
        n = WL[seq[i]]; a.axvspan(x0, x1, color=BAND[n], lw=0, zorder=0)
        if i: a.axvline(x0, color='#404040', lw=1.2, ls=(0, (4, 3)), zorder=1)
        if labels and len(seq) <= 3: a.text((x0 + x1) / 2, 1.01, n, transform=a.get_xaxis_transform(), ha='center', va='bottom', fontsize=9.5, weight='bold', color=INK)
        elif labels: a.text((x0 + x1) / 2, 1.01, seq[i], transform=a.get_xaxis_transform(), ha='center', va='bottom', fontsize=8, color=INK2)

# ---------------- Fig.6
D = {}
f6, ax = plt.subplots(2, 1, figsize=(15, 10.5))
for a, (title, pols, seq, mins, unit, xl) in zip(ax, PANELS):
    bands(a, pols[-1][1], seq, unit)
    v7 = total(pols[-1][1])
    for lab, d, c, lw in pols:
        w = binned(d, mins, unit); D[(title, lab)] = w; t = total(d)
        diff = '' if 'v7' in lab else f'  (v7 {100 * (v7 / t - 1):+.1f}%)'
        a.plot([x for x, _ in w], [y for _, y in w], color=c, lw=lw, marker='o', ms=4.5 if 'v7' in lab else 3.5, label=f'{lab}  전체 {t:.3f}{diff}', zorder=3)
    a.set_title(title, loc='left', fontsize=11.5, color=INK, pad=22); a.set_xlabel(xl); a.set_ylabel('WAF (낮을수록 좋음)')
    a.set_xlim(0, (90 if unit == 1 else 10)); a.legend(loc='upper left', fontsize=9, frameon=True, framealpha=0.95)
ax[1].text(0, -0.17, 'F = FIO-Fast, S = FIO-Slow, V = Varmail', transform=ax[1].transAxes, fontsize=8.5, color=INK2)
f6.suptitle('그림 6. iCAT-v7 vs 고정 CAT: GC 부담이 큰 혼합 워크로드에서의 WAF 추이', x=0.06, ha='left', y=0.985, fontsize=13.5, weight='bold', color=INK)
f6.text(0.06, 0.945, '점선 = 워크로드 전환. 범례 괄호 = 그 정책 대비 v7의 차이(음수 = v7이 낮음 = 좋음). ' + NOTE, fontsize=9, color=INK2)
f6.tight_layout(rect=(0, 0, 1, 0.93)); [f6.savefig(FIG / f'fig6_v7.{e}', dpi=160) for e in ('png', 'pdf')]

# ---------------- Fig.7
C = {}
f7, ax = plt.subplots(2, 1, figsize=(15, 9))
for a, (title, pols, seq, mins, unit, xl) in zip(ax, PANELS):
    d = pols[-1][1]; bands(a, d, seq, unit); c = candidates(d, unit); C[title] = c
    a.step([x for x, _, _ in c], [n for _, n, _ in c], where='post', color=NAVY, lw=2.2, zorder=3)
    resets = [x for (x, n, _), (_, p, _) in zip(c[1:], c) if n > p + 5]
    for x in resets: a.axvline(x, color=DPUR, lw=2, zorder=2)
    a.set_title(title.split(',')[0] + f' — 초기화(보라) {len(resets)}번', loc='left', fontsize=11.5, color=INK, pad=22)
    a.set_ylim(0, 48); a.set_yticks([0, 5, 15, 30, 45]); a.set_xlim(0, (90 if unit == 1 else 10)); a.set_xlabel(xl); a.set_ylabel('남은 후보 수 (파티션 2)')
f7.suptitle('그림 7. iCAT-v7의 학습: 후보 45개 → 나쁜 후보 버리기 → (워크로드 변화 감지) 초기화 → 다시 45개', x=0.06, ha='left', y=0.985, fontsize=13.5, weight='bold', color=INK)
f7.text(0.06, 0.945, '시험 1번 = 후보 하나를 GC 1,024번 동안 써 보고 WAF 점수를 매기는 것. 30초마다 기록. ' + NOTE, fontsize=9, color=INK2)
f7.tight_layout(rect=(0, 0, 1, 0.93)); [f7.savefig(FIG / f'fig7_v7.{e}', dpi=160) for e in ('png', 'pdf')]

# ---------------- xlsx (same data, editable Excel charts; dark paper colors, every line and border 2 pt)
W2, HDR = 25400, PatternFill('solid', fgColor='EEF2FB')
def put(ws, r, c, v, bold=False):
    x = ws.cell(row=r, column=c, value=v)
    if bold: x.font = Font(bold=True); x.fill = HDR; x.alignment = Alignment(wrap_text=True, vertical='center')
    elif isinstance(v, float): x.number_format = '0.000'
def chart(title, xt, yt, xmax, ymin, ymax):
    ch = ScatterChart(); ch.title = title; ch.style = 13; ch.scatterStyle = 'lineMarker'; ch.height, ch.width = 13, 32; ch.legend.position = 'b'
    ch.x_axis.title = xt; ch.y_axis.title = yt; ch.x_axis.delete = False; ch.y_axis.delete = False; ch.x_axis.majorGridlines = None
    ch.x_axis.scaling.min, ch.x_axis.scaling.max, ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, xmax, ymin, ymax
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(solidFill='1A1A1A', w=W2)); return ch
def line(ch, ws, xc, yc, n, color, marker=True, dash=False):
    s = Series(Reference(ws, min_col=yc, min_row=1, max_row=n + 1), Reference(ws, min_col=xc, min_row=2, max_row=n + 1), title_from_data=True)
    s.smooth = False; s.graphicalProperties.line.solidFill = color; s.graphicalProperties.line.width = W2
    if dash: s.graphicalProperties.line.dashStyle = 'dash'
    s.marker.symbol = 'circle' if marker else 'none'
    if marker: s.marker.size = 5; s.marker.graphicalProperties.solidFill = color; s.marker.graphicalProperties.line.solidFill = color
    ch.series.append(s)
def switches(ws, ch, c0, d, seq, unit, ymin, ymax, color='404040', dash=True, xs=None):   # vertical lines as 2-point series
    st = xs if xs is not None else [(s - starts(d)[0]) / 60 / unit for s in starts(d)][1:]
    for i, x in enumerate(st):
        c = c0 + 2 * i; nm = f'→ {WL[seq[i + 1]]}' if xs is None else '초기화'
        put(ws, 1, c, f'선{i + 1} x', True); put(ws, 1, c + 1, nm, True)
        for r, y in ((2, ymin), (3, ymax)): put(ws, r, c, round(x, 3)); put(ws, r, c + 1, y)
        line(ch, ws, c, c + 1, 2, color, marker=False, dash=dash)
    return c0 + 2 * len(st)

wb = Workbook(); wb.remove(wb.active)
for k, (title, pols, seq, mins, unit, xl) in enumerate(PANELS):
    ws = wb.create_sheet(f'그림6{"가나"[k]}'); d = pols[-1][1]
    for i, (lab, dd, c, _) in enumerate(pols):
        w = D[(title, lab)]; put(ws, 1, 2 * i + 1, f'{lab} 시간', True); put(ws, 1, 2 * i + 2, f'{lab} 전체 {total(dd):.3f}', True)
        for r, (x, y) in enumerate(w, 2): put(ws, r, 2 * i + 1, round(x, 3)); put(ws, r, 2 * i + 2, round(y, 4))
    ymax = round(max(y for lab, *_ in pols for _, y in D[(title, lab)]) + 0.5)
    ch = chart(f'그림 6{"가나"[k]}. ' + title.split(' ', 1)[1], xl, 'WAF (낮을수록 좋음)', 90 if unit == 1 else 10, 1, ymax)
    for i, (lab, _, c, _) in enumerate(pols): line(ch, ws, 2 * i + 1, 2 * i + 2, len(D[(title, lab)]), c[1:])
    end = switches(ws, ch, 2 * len(pols) + 2, d, seq, unit, 1, ymax)
    ws.add_chart(ch, f'{ws.cell(row=1, column=end + 1).column_letter}2')
ws = wb.create_sheet('구간별 WAF')   # phase table: one row per phase, one column per policy
for k, (title, pols, seq, *_ ) in enumerate(PANELS):
    r0 = 1 if k == 0 else 7; put(ws, r0, 1, title.split(':')[0], True); put(ws, r0, 2, '워크로드', True)
    for j, (lab, d, *_) in enumerate(pols, 3): put(ws, r0, j, lab, True)
    ph = {lab: phase_waf(d, len(seq)) for lab, d, *_ in pols}
    for i, s in enumerate(seq):
        put(ws, r0 + 1 + i, 1, f'S{i + 1:02d}'); put(ws, r0 + 1 + i, 2, WL[s])
        for j, (lab, *_) in enumerate(pols, 3): put(ws, r0 + 1 + i, j, round(ph[lab][i], 3) if i < len(ph[lab]) else None)
    put(ws, r0 + 1 + len(seq), 1, '전체', True)
    for j, (lab, d, *_) in enumerate(pols, 3): put(ws, r0 + 1 + len(seq), j, round(total(d), 3))
for c, wdt in zip('ABCDEF', [16, 12, 15, 15, 15, 15]): ws.column_dimensions[c].width = wdt
put(ws, 30, 1, NOTE)
wb.save(FIG / 'data' / 'fig6_v7.xlsx')

wb = Workbook(); wb.remove(wb.active)
for k, (title, pols, seq, mins, unit, xl) in enumerate(PANELS):
    ws = wb.create_sheet(f'그림7{"가나"[k]}'); d = pols[-1][1]; c = C[title]
    for j, h in enumerate(['시간', '남은 후보 수 (파티션 2)', '그때 1등 후보(arm)'], 1): put(ws, 1, j, h, True)
    for r, (x, n, b) in enumerate(c, 2): put(ws, r, 1, round(x, 3)); put(ws, r, 2, n); put(ws, r, 3, b)
    ch = chart(f'그림 7{"가나"[k]}. v7 남은 후보 수 — ' + title.split(':')[0][4:], xl, '남은 후보 수', 90 if unit == 1 else 10, 0, 48)
    line(ch, ws, 1, 2, len(c), NAVY[1:], marker=False)
    resets = [x for (x, n, _), (_, p, _) in zip(c[1:], c) if n > p + 5]
    end = switches(ws, ch, 5, d, seq, unit, 0, 48, color=DPUR[1:], dash=False, xs=resets)
    ws.add_chart(ch, f'{ws.cell(row=1, column=end + 1).column_letter}2')
wb.save(FIG / 'data' / 'fig7_v7.xlsx')
print('totals A', {l: total(d) for l, d, *_ in A}, '| B', {l: total(d) for l, d, *_ in B})
print('resets', {t.split(':')[0]: sum(1 for (x, n, _), (_, p, _) in zip(C[t][1:], C[t]) if n > p + 5) for t in C})
