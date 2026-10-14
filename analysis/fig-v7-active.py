"""iCAT-v7 mixY: 후보 수(active) 감소 곡선 = 학습 곡선. 파티션 4개, 1분 간격, 선/테두리 2pt.
-> figs/fig8_v7_active.(png|pdf), figs/data/fig8_v7_active.xlsx"""
import re, datetime as dt
from pathlib import Path
import matplotlib; matplotlib.use('Agg'); import matplotlib.pyplot as plt
from matplotlib import font_manager
font_manager.fontManager.addfont('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
plt.rcParams.update({'font.family': 'Noto Sans CJK JP', 'axes.spines.top': False, 'axes.spines.right': False})
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Font, PatternFill
ROOT = Path(__file__).resolve().parent.parent; FIG = ROOT / 'figs'
R = ROOT / 'result/v7-20261011/mixY/mixY-onlinev7-rep1'
seq = (ROOT / 'result/gc-20261010/mixy-seq.txt').read_text().split()
WL = {'F': 'FIO-Fast', 'S': 'FIO-Slow', 'V': 'Varmail'}
ts = lambda f: dt.datetime.fromisoformat(f.read_text().strip()).timestamp()
t0 = ts(R / 'phase-S01-start.time')
sw = [(ts(R / f'phase-S{i:02d}-start.time') - t0) / 3600 for i in range(2, 21)]
rx = re.compile(r'\[\s*([\d.]+)\] NVMeVirt: WATGC_V2 sample .*?part=(\d) .*?active=(\d+)')
raw = {p: {} for p in range(4)}; k0 = None
for ln in open(R / 'kernel.log'):
    m = rx.search(ln)
    if not m: continue
    k = float(m[1]); k0 = k if k0 is None else k0
    raw[int(m[2])][int((k - k0) // 60)] = int(m[3])        # ponytail: last sample of each minute
mins = sorted(set().union(*raw.values())); last = {p: 45 for p in range(4)}; rows = []
for mi in mins:
    for p in range(4): last[p] = raw[p].get(mi, last[p])
    rows.append([round(mi / 60, 4)] + [last[p] for p in range(4)])
COL = ['1F4E9A', 'B85C00', '1B7A3E', '6B2E8F']
# png/pdf
fig, ax = plt.subplots(figsize=(13, 4.5))
for x in sw: ax.axvline(x, color='#404040', ls='--', lw=1)
for i, w in enumerate(seq): ax.text(i * .5 + .25, 47, w, ha='center', fontsize=9)
for p in range(4): ax.step([r[0] for r in rows], [r[p + 1] for r in rows], where='post', color='#' + COL[p], lw=2, label=f'파티션 {p}')
ax.set_xlabel('시간 (h)'); ax.set_ylabel('남은 후보 수 (active)'); ax.set_ylim(0, 50); ax.set_xlim(0, 10)
for s in ax.spines.values(): s.set_linewidth(2)
ax.legend(loc='upper center', bbox_to_anchor=(.5, -.15), ncol=4, fontsize=9, frameon=False); ax.set_title('iCAT-v7 학습 곡선: 남은 후보 수 (mixY, F/S/V 30분씩)')
fig.tight_layout()
for e in ('png', 'pdf'): fig.savefig(FIG / f'fig8_v7_active.{e}', dpi=150)
# xlsx
W2 = 25400; wb = Workbook(); ws = wb.active; ws.title = 'active'
for c, h in enumerate(['시간(h)'] + [f'파티션 {p}' for p in range(4)], 1):
    x = ws.cell(row=1, column=c, value=h); x.font = Font(bold=True); x.fill = PatternFill('solid', fgColor='EEF2FB')
for r in rows: ws.append(r)
n = len(rows); c0 = 7
for i, x in enumerate(sw):                                   # 부하 전환 = 세로 점선 (2점 series)
    c = c0 + 2 * i; ws.cell(1, c, 'x'); ws.cell(1, c + 1, f'→ {WL[seq[i + 1]]}')
    ws.cell(2, c, x); ws.cell(3, c, x); ws.cell(2, c + 1, 0); ws.cell(3, c + 1, 50)
ch = ScatterChart(); ch.title = 'iCAT-v7 학습 곡선: 남은 후보 수 (mixY)'; ch.style = 13; ch.height, ch.width = 13, 32
ch.x_axis.title = '시간 (h)'; ch.y_axis.title = '남은 후보 수 (active)'; ch.x_axis.delete = ch.y_axis.delete = False
ch.x_axis.majorGridlines = None; ch.legend.position = 'b'
ch.x_axis.scaling.min, ch.x_axis.scaling.max, ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 10, 0, 50
ch.graphical_properties = GraphicalProperties(ln=LineProperties(solidFill='1A1A1A', w=W2))
for p in range(4):
    s = Series(Reference(ws, min_col=p + 2, min_row=1, max_row=n + 1), Reference(ws, min_col=1, min_row=2, max_row=n + 1), title_from_data=True)
    s.smooth = False; s.marker.symbol = 'none'; s.graphicalProperties.line.solidFill = COL[p]; s.graphicalProperties.line.width = W2
    ch.series.append(s)
for i in range(len(sw)):
    c = c0 + 2 * i
    s = Series(Reference(ws, min_col=c + 1, min_row=2, max_row=3), Reference(ws, min_col=c, min_row=2, max_row=3))
    s.marker.symbol = 'none'; s.graphicalProperties.line.solidFill = '404040'; s.graphicalProperties.line.width = W2; s.graphicalProperties.line.dashStyle = 'dash'
    ch.series.append(s)
ws.add_chart(ch, 'H25'); wb.save(FIG / 'data' / 'fig8_v7_active.xlsx')
print(n, 'rows')
