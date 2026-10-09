#!/usr/bin/env python3
"""mixU (FIO-Fast <-> Varmail 번갈아 8구간, 각 4500 s = 10시간), iCAT-v4 1회 (두 번째 PC, 2026-10-05~06).
그림 6형: 20분 구간 WAF 추이(점 표시) + 구간별 WAF. 그림 7형: 학습 회차별 남은 후보 수(회차를 같은 폭으로 이어 그림).
-> figs/fig6_mixU_v4.png, figs/fig7_mixU_v4.png, figs/data/mixU_v4.xlsx
원천: result/mixU-onlinev4-after-v6-20261005/mixU-onlinev4-rep1 (control-series, phase-X-start.time, summary, kernel 로그)."""
import re, datetime as dt, importlib.util
from pathlib import Path
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, BarChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.legend import LegendEntry
from openpyxl.styles import Font, PatternFill, Alignment
font_manager.fontManager.addfont('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
spec = importlib.util.spec_from_file_location('v4log', '/home/oy/iCAT/analysis/v4-log.py'); v4log = importlib.util.module_from_spec(spec); spec.loader.exec_module(v4log)
D = Path('/home/oy/iCAT/result/mixU-onlinev4-after-v6-20261005/mixU-onlinev4-rep1'); FIG = Path('/home/oy/iCAT/figs')
PART, STEP = '2', 20 / 60
SL = 'ABCDEFGH'; NAMES = ['FIO-Fast', 'Varmail'] * 4
TYPE = {'FIO-Fast': '합성 덮어쓰기 (FIO 3영역, 6 GiB)', 'Varmail': '메일 서버 흉내 (Filebench)'}
BAND = {'FIO-Fast': '#f7ecec', 'Varmail': '#eef2fb'}
INK, INK2, BLUE, PUR = '#1d1d1f', '#6e6e73', '#2a78d6', '#7d3cff'
plt.rcParams.update({'font.family': 'Noto Sans CJK JP', 'font.size': 10.5, 'axes.spines.top': False, 'axes.spines.right': False,
                     'axes.spines.left': False, 'axes.edgecolor': '#d2d2d7', 'axes.grid': True, 'axes.grid.axis': 'y', 'grid.color': '#ededf0',
                     'axes.axisbelow': True, 'xtick.color': INK2, 'ytick.color': INK2, 'ytick.left': False, 'axes.labelcolor': INK2})

raw = [tuple(map(int, m.groups())) for l in open(D / 'control-series.txt') if (m := re.match(r'(\d+) .*active=1 .*host_pages=(\d+) gc_pages=(\d+)', l))]
T0 = raw[0][0]; s = [((t - T0) / 3600, h, g) for t, h, g in raw]
W, k = [], 0.0
while k + STEP <= s[-1][0] + 1e-9:
    a = min(s, key=lambda p: abs(p[0] - k)); b = min(s, key=lambda p: abs(p[0] - k - STEP))
    if b[1] > a[1]: W.append((round(k + STEP / 2, 4), round(1 + (b[2] - a[2]) / (b[1] - a[1]), 4)))
    k += STEP
X1 = s[-1][0]
starts = [(dt.datetime.fromisoformat((D / f'phase-{c}-start.time').read_text().strip()).timestamp() - T0) / 3600 for c in SL]
bounds = starts + [X1]; mids = [(bounds[i] + bounds[i + 1]) / 2 for i in range(8)]
sm = (D / 'summary.txt').read_text()
TOT = float(re.search(r'^total .*WAF=([\d.]+)', sm, re.M)[1])
PH = [float(re.search(rf'^phase{c}\(\S+\) .*WAF=([\d.]+)', sm, re.M)[1]) for c in SL]

L = v4log.lines(str(D))
ts = lambda l: float(re.match(r'\[\s*(\d+\.\d+)\]', l)[1])
t0 = next(ts(l) for l in L if 'Measurement: command=start' in l)
tr = [((ts(l) - t0) / 3600, int(m[1]), int(m[2])) for l in L if (m := re.search(rf'part={PART} epoch=(\d+) window=\d+ .*?active=(\d+)', l))]
resets = [(ts(l) - t0) / 3600 for l in L if 'WATGC_V2 reset' in l and f'part={PART} ' in l]
best = {}
for l in L:
    if (m := re.search(rf'part={PART} epoch=(\d+) window=\d+ .*?best=(\d+) ', l)): best[int(m[1])] = int(m[2])
rounds = []
for ep in sorted({e for _, e, _ in tr}):
    v = [(h, a) for h, e, a in tr if e == ep]; h0, h1 = v[0][0], v[-1][0]
    wl = [f'{i + 1}.{n}' for i, n in enumerate(NAMES) if min(h1, bounds[i + 1]) - max(h0, bounds[i]) > 0.25]
    rounds.append((ep, h0, h1, wl, [a for _, a in v]))
RW = {i: f'{r[0]}회차\n{"+".join(w.split(".")[1] for w in r[3]) or "-"}\n{r[1]:.1f}~{r[2]:.1f}h · 시험 {len(r[4])}번' for i, r in enumerate(rounds)}
CH = [(i + j / len(r[4]), v) for i, r in enumerate(rounds) for j, v in enumerate(r[4]) if j == 0 or v != r[4][j - 1]]
SEQ = []
for i, r in enumerate(rounds):
    n = len(r[4]); keep = set(range(0, n, max(1, n // 70))) | {n - 1}; pts = []
    for j, v in enumerate(r[4]):
        x = i + j / n
        if j and v != r[4][j - 1]: pts.append((x, r[4][j - 1]))
        if j in keep or (j and v != r[4][j - 1]): pts.append((x, v))
    SEQ.append(pts)
NOTE = 'iCAT-v4 1회(두 번째 PC). 구간 8개 × 4500초. 고정 CAT(기본·견고·최적)은 아직 이 혼합으로 측정하지 않음.'

# ---- PNG
f6, a6 = plt.subplots(figsize=(14, 5.4))
for i, n in enumerate(NAMES):
    a6.axvspan(bounds[i], bounds[i + 1], color=BAND[n], lw=0, zorder=0)
    if i: a6.axvline(bounds[i], color='#5a5955', lw=1.1, ls=(0, (4, 3)), zorder=1)
    a6.text(mids[i], 1.02, f'{i + 1}. {n}\nWAF {PH[i]:.2f}', transform=a6.get_xaxis_transform(), ha='center', va='bottom', fontsize=9.5, color=INK)
a6.plot([h for h, _ in W], [v for _, v in W], color=BLUE, lw=2.4, marker='o', ms=4.5, label=f'iCAT-v4   10시간 전체 {TOT:.3f}')
for r in resets: a6.axvline(r, color=PUR, lw=1.4, alpha=0.8, zorder=2)
a6.set_xlim(0, X1); a6.set_xticks(range(0, int(X1) + 1)); a6.set_xlabel('측정 시작 후 시간 (시간)'); a6.set_ylabel('구간 WAF (20분, 낮을수록 좋음)')
a6.legend(loc='upper center', bbox_to_anchor=(0.5, -0.13), frameon=False)
f6.subplots_adjust(top=0.8)
f6.suptitle('mixU (FIO-Fast ↔ Varmail 번갈아 10시간): iCAT-v4의 WAF 추이', x=0.06, ha='left', y=0.99, fontsize=13.5, weight='bold', color=INK)
f6.text(0.06, 0.935, '점 = 20분 구간. 점선 = 워크로드 전환, 보라 선 = v4가 변화를 감지해 다시 학습한 시점. ' + NOTE, fontsize=9, color=INK2)
f6.savefig(FIG / 'fig6_mixU_v4.png', dpi=200, bbox_inches='tight')

f7, a7 = plt.subplots(figsize=(14, 6.2))
for i, r in enumerate(rounds):
    wl = r[3][0].split('.')[1] if r[3] else 'Varmail'
    a7.axvspan(i, i + 1, color=BAND[wl], lw=0, zorder=0)
    a7.text(i + 0.5, -0.04, RW[i], transform=a7.get_xaxis_transform(), ha='center', va='top', fontsize=8.5, color=INK, linespacing=1.3)
    if i: a7.axvline(i, color=PUR, lw=2, zorder=4)
a7.plot([x for p in SEQ for x, _ in p], [y for p in SEQ for _, y in p], color=BLUE, lw=2.2, zorder=3)
a7.scatter([x for x, _ in CH], [y for _, y in CH], s=26, color=BLUE, edgecolor='white', linewidth=0.8, zorder=5)
for i, r in enumerate(rounds):
    lo = min(r[4]); jl = r[4].index(lo)
    if lo < r[4][0]: a7.annotate(f'{r[4][0]}→{lo}개\n최고 후보 {best.get(r[0], "-")}번', (i + jl / len(r[4]), lo), (5, 6), textcoords='offset points', fontsize=9, color=BLUE, weight='bold')
a7.set_xlim(0, len(rounds)); a7.set_xticks([]); a7.set_ylim(0, 50); a7.set_yticks([0, 5, 15, 30, 45]); a7.set_ylabel(f'남은 후보 수 (파티션 {PART})'); a7.grid(False); a7.yaxis.grid(True)
f7.subplots_adjust(top=0.82, bottom=0.22)
f7.suptitle('mixU: iCAT-v4의 학습 — 후보 줄이기 → (전환 감지) 초기화 → 다시 줄이기', x=0.06, ha='left', y=0.99, fontsize=13.5, weight='bold', color=INK)
f7.text(0.06, 0.92, '칸 하나 = 학습 1회차(폭 같게). 보라 선 = 초기화(v4는 45개 전부가 아니라 살아남은 후보+이웃만 되살림). 점 = 회차 시작과 후보를 버린 순간.\n 시험 1번 = 후보 하나를 512 MiB(GC 128번 이상) 써 보고 점수 매기기.', fontsize=9, color=INK2)
f7.savefig(FIG / 'fig7_mixU_v4.png', dpi=200, bbox_inches='tight')

# ---- xlsx
HDR = PatternFill('solid', fgColor='EEF2FB')
def put(ws, r, c, v, bold=False):
    x = ws.cell(row=r, column=c, value=v)
    if bold: x.font = Font(bold=True); x.fill = HDR; x.alignment = Alignment(wrap_text=True, vertical='center')
def scatter(title, xt, yt, xmin, xmax, unit, ymin, ymax, h=12, w=32):
    ch = ScatterChart(); ch.title = title; ch.style = 13; ch.scatterStyle = 'lineMarker'; ch.height, ch.width = h, w; ch.legend.position = 'b'
    ch.x_axis.title = xt; ch.y_axis.title = yt; ch.x_axis.delete = False; ch.y_axis.delete = False
    ch.x_axis.scaling.min, ch.x_axis.scaling.max, ch.x_axis.majorUnit = xmin, xmax, unit; ch.y_axis.scaling.min, ch.y_axis.scaling.max = ymin, ymax
    ch.x_axis.majorGridlines = None; return ch
def series(ws, xc, yc, n, color, width, marker, size=4, r0=2, noline=False):
    s = Series(Reference(ws, min_col=yc, min_row=1, max_row=r0 + n - 1), Reference(ws, min_col=xc, min_row=r0, max_row=r0 + n - 1), title_from_data=True)
    s.smooth = False; s.graphicalProperties.line.solidFill = color; s.graphicalProperties.line.width = width
    if noline: s.graphicalProperties.line.solidFill = None; s.graphicalProperties.line.noFill = True   # 10-09: both fills in one <a:ln> made Excel drop the chart
    s.marker.symbol = marker
    if marker != 'none': s.marker.size = size; s.marker.graphicalProperties.solidFill = color; s.marker.graphicalProperties.line.solidFill = 'FFFFFF' if noline else color
    return s
def deco(ch, ws, c0, xs, ytop, ybot, labels, color='8A8984', dash=True):
    """10-07: same construction as the earlier working Excel figures (fig7-mixT.py): one 2-point series per vertical line,
    named after the segment that starts there, shown in the legend. No data-label series, no hidden legend entries."""
    for i, x in enumerate(xs):
        c = c0 + 2 * i; name = '→ ' + labels[i + 1][1] if i + 1 < len(labels) else f'선{i + 1}'
        put(ws, 1, c, f'{name} x', True); put(ws, 1, c + 1, name, True)
        for r, y in ((2, ybot), (3, ytop)): put(ws, r, c, round(x, 3)); put(ws, r, c + 1, y)
        s = series(ws, c, c + 1, 2, color, 12700, 'none')
        if dash: s.graphicalProperties.line.dashStyle = 'dash'
        ch.series.append(s)
    return c0 + 2 * len(xs)

wb = Workbook(); ws = wb.active; ws.title = 'WAF 추이'
put(ws, 1, 1, '시간(h)', True); put(ws, 1, 2, f'iCAT-v4 (10시간 전체 {TOT:.3f})', True)
for i, (h, v) in enumerate(W, 2): put(ws, i, 1, h); put(ws, i, 2, v)
ymax = round(max(v for _, v in W) + 0.8)
ch = scatter('mixU: 20분 구간 WAF 추이 (점 = 20분, 낮을수록 좋음)', '측정 시작 후 시간 (시간)', 'WAF', 0, 10, 1, 0, ymax, h=13, w=34)
ch.series.append(series(ws, 1, 2, len(W), '2A78D6', 31750, 'circle', 6))
end = deco(ch, ws, 4, bounds[1:8], ymax - 0.3, 0, [(mids[i], f'{i + 1}.{n} WAF {PH[i]:.2f}') for i, n in enumerate(NAMES)])
ws.add_chart(ch, f'{ws.cell(row=1, column=end + 1).column_letter}2')
put(ws, 30, end + 1, '점선 = 워크로드 전환, 위 글자 = 구간 이름과 그 구간 평균 WAF. ' + NOTE)

w2 = wb.create_sheet('구간별 WAF')
for j, h in enumerate(['구간', '워크로드', '종류', 'WAF (iCAT-v4)', '시작(h)'], 1): put(w2, 1, j, h, True)
for i, n in enumerate(NAMES):
    for j, v in enumerate([f'{i + 1}. {n}', n, TYPE[n], round(PH[i], 3), round(starts[i], 2)], 1): put(w2, 2 + i, j, v)
for j, v in enumerate(['10시간 전체', '', '', round(TOT, 3), ''], 1): put(w2, 10, j, v, True)
for j, wd in enumerate([13, 10, 30, 14, 9], 1): w2.column_dimensions[w2.cell(row=1, column=j).column_letter].width = wd
bc = BarChart(); bc.type = 'col'; bc.title = '구간별 평균 WAF (iCAT-v4, 낮을수록 좋음)'; bc.y_axis.title = 'WAF'; bc.height, bc.width = 10, 24
bc.add_data(Reference(w2, min_col=4, min_row=1, max_row=9), titles_from_data=True); bc.set_categories(Reference(w2, min_col=1, min_row=2, max_row=9))
bc.series[0].graphicalProperties.solidFill = '2A78D6'; bc.legend = None; bc.x_axis.delete = False; bc.y_axis.delete = False
w2.add_chart(bc, 'A13')

w3 = wb.create_sheet('학습 (후보 수)')
xs = [round(x, 4) for p in SEQ for x, _ in p]; ys = [y for p in SEQ for _, y in p]
put(w3, 1, 1, '진행 위치 (회차 + 회차 안 비율)', True); put(w3, 1, 2, f'남은 후보 수 (파티션 {PART})', True)
for i, (x, y) in enumerate(zip(xs, ys), 2): put(w3, i, 1, x); put(w3, i, 2, y)
put(w3, 1, 3, '점 x', True); put(w3, 1, 4, '회차 시작·후보 버린 순간', True)
for i, (x, y) in enumerate(CH, 2): put(w3, i, 3, round(x, 4)); put(w3, i, 4, y)
ch3 = scatter('mixU: v4 학습 — 줄이기 → 초기화(보라, 살아남은 후보+이웃) → 다시 (칸 하나 = 학습 1회차)', '학습 회차 (아래 표: 회차별 워크로드·시간)', '남은 후보 수', 0, len(rounds), 1, 0, 50, h=13, w=34)
ch3.series.append(series(w3, 1, 2, len(xs), '2A78D6', 25400, 'none'))
ch3.series.append(series(w3, 3, 4, len(CH), '2A78D6', 0, 'circle', 6, noline=True))
end3 = deco(ch3, w3, 6, list(range(1, len(rounds))), 48, 0, [(i + 0.5, RW[i].split('\n')[0] + ' ' + RW[i].split('\n')[1]) for i in range(len(rounds))], '7D3CFF', dash=False)
w3.add_chart(ch3, f'{w3.cell(row=1, column=end3 + 1).column_letter}2')
r0 = 32
for j, h in enumerate(['회차', '워크로드', '시작(h)', '끝(h)', '시험 횟수', '처음 후보', '최소 후보', '마지막 최고 후보'], 1): put(w3, r0, end3 + j, h, True)
for i, r in enumerate(rounds):
    for j, v in enumerate([r[0], '+'.join(r[3]), round(r[1], 2), round(r[2], 2), len(r[4]), r[4][0], min(r[4]), best.get(r[0])], 1): put(w3, r0 + 1 + i, end3 + j, v)
w4 = wb.create_sheet('판단별 원자료')
for j, h in enumerate(['시간(h)', '학습 회차', f'남은 후보 수 (파티션 {PART})'], 1): put(w4, 1, j, h, True)
for i, (h, e, a) in enumerate(tr, 2): put(w4, i, 1, round(h, 4)); put(w4, i, 2, e); put(w4, i, 3, a)
wb.save(FIG / 'data' / 'mixU_v4.xlsx')
# Fig.7 file on its own (same layout as fig7_v6_mixT.xlsx): only the learning sheets
for name in ('WAF 추이', '구간별 WAF'): wb.remove(wb[name])
wb.save(FIG / 'data' / 'fig7_mixU_v4.xlsx')
print('total', TOT, 'phases', PH, 'resets', [round(r, 2) for r in resets])
for r in rounds: print(r[0], r[3], len(r[4]), r[4][0], '->', min(r[4]), 'best', best.get(r[0]))
