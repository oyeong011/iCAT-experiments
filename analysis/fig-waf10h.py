#!/usr/bin/env python3
"""10시간 실험 종합 그림: (가) 워크로드별 10시간 전체 WAF — iCAT-v4 vs 최적(47) / 견고(50) / 기본(37),
(나) 불규칙 순서 9구간의 구간별 WAF. -> figs/fig_waf10h.(png|svg|pdf), figs/data/fig_waf10h.xlsx (막대 차트 포함)
원천: 각 실행의 summary.txt (각 1회 측정; CAT-37/50 일부는 두 번째 PC)."""
import re
from pathlib import Path
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font
font_manager.fontManager.addfont('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
R = Path('/home/oy/iCAT/result'); M = R / 'mix-20260911'
POL = [('iCAT-v4 (제안)', '#2a78d6'), ('CAT-47 (최적)', '#2b2a28'), ('CAT-50 (견고)', '#9a9994'), ('CAT-37 (기본)', '#e3a33b')]
SETS = [('불규칙 순서 9구간', {'iCAT-v4 (제안)': M / 'mixT-onlinev4-rep1', 'CAT-47 (최적)': M / 'mixT-fixed47-rep1',
                            'CAT-50 (견고)': R / 'mixT-fixed50-tenhour-20261004/mixT-fixed50-rep1', 'CAT-37 (기본)': R / 'mixT-fixed37-tenhour-20261004/mixT-fixed37-rep1'}),
        ('OLTP → Varmail', {'iCAT-v4 (제안)': M / 'mixO-onlinev4-rep1-x20', 'CAT-50 (견고)': M / 'mixO-fixed50-rep1-x20',
                            'CAT-37 (기본)': R / 'mixO-fixed37-chunked-20261003/mixO-fixed37-rep1-x20'}),
        ('FIO-Fast → Varmail', {'iCAT-v4 (제안)': M / 'mixF-onlinev4-rep1-x30', 'CAT-50 (견고)': R / 'mixF-fixed50-tenhour-20261004/mixF-fixed50-rep1-x30',
                                'CAT-37 (기본)': R / 'mixF-fixed37-tenhour-20261004/mixF-fixed37-rep1-x30'})]
NM = {'sqlite-a': 'YCSB-A', 'sqlite-b': 'YCSB-B', 'oltp': 'OLTP', 'varmail': 'Varmail', 'test4': 'FIO-Fast'}
def waf(d):
    s = (d / 'summary.txt').read_text()
    return float(re.search(r'^total .*WAF=([\d.]+)', s, re.M)[1]), [(NM.get(m[1], m[1]), float(m[2])) for m in re.finditer(r'^phase[A-I]\((\S+)\) .*WAF=([\d.]+)', s, re.M)]
T = {w: {p: waf(d) for p, d in runs.items() if (d / 'summary.txt').exists()} for w, runs in SETS}
INK, INK2 = '#1d1d1f', '#6e6e73'
plt.rcParams.update({'font.family': 'Noto Sans CJK JP', 'font.size': 10.5, 'axes.facecolor': 'white', 'figure.facecolor': 'white',
                     'axes.spines.top': False, 'axes.spines.right': False, 'axes.spines.left': False, 'axes.edgecolor': '#d2d2d7',
                     'axes.grid': True, 'axes.grid.axis': 'y', 'grid.color': '#ededf0', 'axes.axisbelow': True,
                     'xtick.color': INK, 'ytick.color': INK2, 'ytick.left': False, 'svg.fonttype': 'path'})
fig, (a1, a2) = plt.subplots(2, 1, figsize=(13, 9.5), gridspec_kw={'height_ratios': [1, 1.15], 'hspace': 0.42})
W = 0.2
for i, (w, _) in enumerate(SETS):
    v4 = T[w]['iCAT-v4 (제안)'][0]
    for j, (p, c) in enumerate(POL):
        x = i + (j - 1.5) * W
        if p not in T[w]:
            a1.text(x, 1.02, '미측정', ha='center', va='bottom', fontsize=8, color=INK2, rotation=90); continue
        v = T[w][p][0]
        a1.bar(x, v - 1, bottom=1, width=W * 0.9, color=c, label=p if i == 0 else None)
        a1.text(x, v + 0.02, f'{v:.3f}', ha='center', va='bottom', fontsize=9, color=INK, weight='bold' if 'v4' in p else 'normal')
        if 'v4' not in p:
            d = (v4 / v - 1) * 100
            a1.text(x, 1.04, f'v4\n{d:+.1f}%', ha='center', va='bottom', fontsize=8, color='white' if p != 'CAT-50 (견고)' else INK, linespacing=1.1)
a1.set_xticks(range(len(SETS))); a1.set_xticklabels([w for w, _ in SETS], fontsize=11, weight='bold')
a1.set_ylim(1, 2.75); a1.set_ylabel('10시간 전체 WAF (낮을수록 좋음)', color=INK2)
a1.legend(loc='upper left', ncol=4, frameon=False, fontsize=9.5)
a1.set_title('(가) 10시간 실험 3종의 전체 WAF — 막대 안 % = 그 정책 대비 iCAT-v4의 차이(음수 = v4가 낮음)', loc='left', fontsize=11, weight='bold', color=INK)
w = '불규칙 순서 9구간'; names = [n for n, _ in T[w]['iCAT-v4 (제안)'][1]]
for j, (p, c) in enumerate(POL):
    ph = T[w][p][1]
    a2.bar([k + (j - 1.5) * W for k in range(len(ph))], [v - 1 for _, v in ph], bottom=1, width=W * 0.9, color=c, label=p)
a2.set_xticks(range(len(names))); a2.set_xticklabels([f'{k + 1}. {n}' for k, n in enumerate(names)], fontsize=9.5)
a2.set_ylabel('구간 WAF (낮을수록 좋음)', color=INK2); a2.legend(loc='upper left', ncol=4, frameon=False, fontsize=9.5)
a2.set_title('(나) 불규칙 순서 9구간의 구간별 WAF (각 4000초)', loc='left', fontsize=11, weight='bold', color=INK)
fig.suptitle('10시간 실험 결과 종합: iCAT-v4 vs 고정 CAT (최적·견고·기본)', x=0.06, ha='left', y=0.995, fontsize=13.5, weight='bold', color=INK)
fig.text(0.06, 0.955, '각 1회 측정. CAT-47(최적)은 불규칙 순서 실험에서만 측정. 일부 고정 CAT은 두 번째 PC에서 측정(CAT 검증값 0.2% 이내 일치).', fontsize=9.5, color=INK2)
for e in ('png', 'svg', 'pdf'): fig.savefig(f'/home/oy/iCAT/figs/fig_waf10h.{e}', dpi=200, bbox_inches='tight')

wb = Workbook(); ws = wb.active; ws.title = '10시간 전체 WAF'
ws.append(['워크로드'] + [p for p, _ in POL]); [setattr(c, 'font', Font(bold=True)) for c in ws[1]]
for w, _ in SETS: ws.append([w] + [round(T[w][p][0], 3) if p in T[w] else None for p, _ in POL])
ws.append([]); ws.append(['iCAT-v4 대비 차이 (%) = v4/그 정책 − 1 (음수 = v4가 낮음)'])
for w, _ in SETS: ws.append([w] + [round((T[w]['iCAT-v4 (제안)'][0] / T[w][p][0] - 1) * 100, 1) if p in T[w] else None for p, _ in POL])
ch = BarChart(); ch.type = 'col'; ch.title = '10시간 전체 WAF (낮을수록 좋음)'; ch.y_axis.title = 'WAF'; ch.height, ch.width = 10, 22
ch.add_data(Reference(ws, min_col=2, max_col=5, min_row=1, max_row=4), titles_from_data=True); ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
ch.y_axis.scaling.min = 1.0; ch.y_axis.delete = False; ch.x_axis.delete = False; ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True
for s, (_, c) in zip(ch.series, POL): s.graphicalProperties.solidFill = c[1:].upper(); s.graphicalProperties.line.solidFill = c[1:].upper()
ws.add_chart(ch, 'H2')
w = '불규칙 순서 9구간'   # w was reused by the loops above
ws2 = wb.create_sheet('불규칙 9구간 구간별')
ws2.append(['구간'] + [p for p, _ in POL]); [setattr(c, 'font', Font(bold=True)) for c in ws2[1]]
for k, n in enumerate(names): ws2.append([f'{k + 1}. {n}'] + [round(T[w][p][1][k][1], 3) for p, _ in POL])
ws2.append(['10시간 전체'] + [round(T[w][p][0], 3) for p, _ in POL])
ch2 = BarChart(); ch2.type = 'col'; ch2.title = '불규칙 순서 9구간: 구간별 WAF'; ch2.y_axis.title = 'WAF'; ch2.height, ch2.width = 11, 30
ch2.add_data(Reference(ws2, min_col=2, max_col=5, min_row=1, max_row=10), titles_from_data=True); ch2.set_categories(Reference(ws2, min_col=1, min_row=2, max_row=10))
ch2.y_axis.scaling.min = 1.0; ch2.y_axis.delete = False; ch2.x_axis.delete = False
for s, (_, c) in zip(ch2.series, POL): s.graphicalProperties.solidFill = c[1:].upper(); s.graphicalProperties.line.solidFill = c[1:].upper()
ws2.add_chart(ch2, 'H2')
wb.save('/home/oy/iCAT/figs/data/fig_waf10h.xlsx'); print('ok')
