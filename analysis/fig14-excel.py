#!/usr/bin/env python3
"""fig14 30초 WAF를 엑셀용으로: 워크로드마다 정책별 '시간 | WAF' 열 쌍 (CSV 4개) + 시트 4개에 꺾은선 차트가 들어간 xlsx
원천: figs/data/fig14_30s.csv. 정책마다 실제 기록 시각을 그대로 둠(분산형 차트)."""
import csv, collections
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
D = Path('/home/oy/iCAT/figs/data')
r = list(csv.DictReader(open(D / 'fig14_30s.csv', encoding='utf-8-sig')))
POL = ['iCAT-v4', 'CAT-47 (최적)', 'CAT-50 (견고)', 'CAT-37 (기본)']; COL = ['2A78D6', '2B2A28', 'A3A29D', 'E3A33B']
wb = Workbook(); wb.remove(wb.active)
for wl in dict.fromkeys(x['워크로드'] for x in r):
    # exact times kept: each policy gets its own (시간, WAF) column pair, so no point is merged or lost
    # dropped (near-empty) intervals are left out entirely, so Excel draws one unbroken line per policy
    cols = {p: [(float(x['경과(분)']), float(x['30초 WAF'])) for x in r if x['워크로드'] == wl and x['정책'] == p and x['30초 WAF']] for p in POL}
    sw = {p: next((float(x['경과(분)']) for x in r if x['워크로드'] == wl and x['정책'] == p and x['구간'].startswith('뒤')), None) for p in POL}
    n = max(map(len, cols.values()))
    head = [h for p in POL for h in (f'{p} 시간(분)', f'{p} WAF')]
    rows = [[v for p in POL for v in (cols[p][i] if i < len(cols[p]) else (None, None))] for i in range(n)]
    name = wl.replace(' → ', '-to-').replace(' ', '')
    with open(D / f'fig14_lines_{name}.csv', 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f); w.writerow(head); w.writerows([['' if v is None else (f'{v:.4f}' if j % 2 else v) for j, v in enumerate(x)] for x in rows])
    ws = wb.create_sheet(wl.replace(' → ', '→')[:31])
    ws.append(head); [ws.append(x) for x in rows]
    ws['J1'] = '워크로드 전환 시점(분)'; [ws.cell(row=i + 2, column=10, value=f'{p}: {sw[p]:.1f}') for i, p in enumerate(POL)]
    ch = ScatterChart(); ch.title = f'{wl}: 30초마다 잰 WAF (낮을수록 좋음)'; ch.style = 13
    ch.x_axis.title = '실험 시작 후 시간 (분)'; ch.y_axis.title = '30초 WAF'; 
    ch.x_axis.delete = False; ch.y_axis.delete = False; ch.height, ch.width = 14, 30; ch.legend.position = 'b'
    ch.scatterStyle = 'lineMarker'; ch.display_blanks = 'span'   # without scatterStyle Excel may draw markers only (looks broken)
    ys = [v for p in POL for _, v in cols[p]]; xm = max(t for p in POL for t, _ in cols[p])
    lo, hi = min(ys), max(ys); pad = (hi - lo) * 0.08
    ch.y_axis.scaling.min = round(lo - pad, 2); ch.y_axis.scaling.max = round(hi + pad, 2)   # Excel would otherwise start at 0 and squash the lines
    ch.x_axis.scaling.min = 0; ch.x_axis.scaling.max = round(xm + 0.5); ch.x_axis.majorUnit = 5
    ch.y_axis.number_format = '0.00'; ch.x_axis.number_format = '0'
    for j, c in enumerate(COL):
        xs = Reference(ws, min_col=2 * j + 1, min_row=2, max_row=n + 1)
        sr = Series(Reference(ws, min_col=2 * j + 2, min_row=1, max_row=n + 1), xs, title_from_data=True)
        sr.marker.symbol = 'none'; sr.smooth = False; sr.graphicalProperties.line.solidFill = c; sr.graphicalProperties.line.width = 32000 if j == 0 else 15000
        ch.series.append(sr)
    ws.add_chart(ch, 'L2'); ws.column_dimensions['J'].width = 26
wb.save(D / 'fig14_lines.xlsx'); print('ok')
